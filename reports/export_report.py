#!/usr/bin/env python3
"""
Excel / PDF exports for the HRIS (Python).

  python3 export_report.py attendance --db hris.db --month 2026-11 --format xlsx|pdf [--today YYYY-MM-DD]
  python3 export_report.py leave-form  --db hris.db --id 5
  python3 export_report.py employees   --db hris.db

The file is written to standard output (binary) so the Node server can stream it to the browser.
Needs:  pip install -r reports/requirements.txt   (openpyxl for Excel, reportlab for PDF)
"""
import argparse, io, json, os, sqlite3, sys
from datetime import date, datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from attendance_report import build_report  # same folder: the monthly calculation lives in one place


def fail(msg, code=1):
    print(json.dumps({"error": msg}), file=sys.stderr)
    sys.exit(code)


def need(module, package):
    try:
        return __import__(module)
    except ImportError:
        fail(f"The Python package '{package}' is not installed. Run:  pip install -r reports/requirements.txt", 2)


def connect(db):
    con = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
    con.row_factory = sqlite3.Row
    return con


# ------------------------------------------------------------------ Excel
def attendance_xlsx(report):
    need("openpyxl", "openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"
    ws["A1"] = f"Attendance Report - {report['month']}"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = f"Counted through {report['through']}  |  Generated {datetime.now():%Y-%m-%d %H:%M}"
    ws["A2"].font = Font(italic=True, color="666666")

    heads = ["Employee No", "Name", "Department", "Shift", "Scheduled Days", "Present", "Late", "Absent",
             "On Leave", "Late Minutes", "Hours Worked", "Attendance Rate %"]
    hr = 4
    fill = PatternFill("solid", fgColor="1F5EFF")
    thin = Side(style="thin", color="D0D7E2")
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, r in enumerate(report["rows"], hr + 1):
        vals = [r["employee_no"], r["name"], r["department"], r["shift"], r["scheduled_days"], r["present"], r["late"],
                r["absent"], r["on_leave"], r["late_minutes"], r["hours_worked"], r["attendance_rate"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.border = Border(bottom=thin)
            if c >= 5:
                cell.alignment = Alignment(horizontal="right")
        ws.cell(row=i, column=12).number_format = "0.0"
        ws.cell(row=i, column=11).number_format = "0.0"
    first, last = hr + 1, hr + len(report["rows"])
    tr = last + 1
    ws.cell(row=tr, column=2, value="TOTAL / AVERAGE").font = Font(bold=True)
    if report["rows"]:
        for c in range(5, 12):
            L = get_column_letter(c)
            ws.cell(row=tr, column=c, value=f"=SUM({L}{first}:{L}{last})").font = Font(bold=True)
        ws.cell(row=tr, column=12, value=f"=IFERROR(AVERAGE(L{first}:L{last}),\"\")").font = Font(bold=True)
        ws.cell(row=tr, column=12).number_format = "0.0"
    for c, w in enumerate([13, 24, 22, 20, 11, 9, 7, 8, 9, 11, 11, 13], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[hr].height = 32
    ws.freeze_panes = ws.cell(row=hr + 1, column=3)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def employees_xlsx(db):
    need("openpyxl", "openpyxl")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    con = connect(db)
    rows = con.execute("""
        SELECT e.employee_no, e.first_name, e.last_name, COALESCE(d.name,'') dept, COALESCE(e.position,'') pos,
               COALESCE(e.email,'') email, COALESCE(e.mobile,'') mobile, COALESCE(e.date_hired,'') hired,
               r.name role, CASE u.is_active WHEN 1 THEN 'Active' ELSE 'Inactive' END status,
               COALESCE(s.name,'') shift
        FROM employees e JOIN users u ON u.id = e.user_id JOIN roles r ON r.id = u.role_id
        LEFT JOIN departments d ON d.id = e.department_id
        LEFT JOIN shift_assignments sa ON sa.employee_id = e.id LEFT JOIN shifts s ON s.id = sa.shift_id
        ORDER BY e.last_name, e.first_name""").fetchall()
    con.close()
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    heads = ["Employee No", "First Name", "Last Name", "Department", "Position", "Email", "Mobile", "Date Hired", "Role", "Status", "Shift"]
    for c, h in enumerate(heads, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F5EFF")
        cell.alignment = Alignment(horizontal="center")
    for i, r in enumerate(rows, 2):
        for c, v in enumerate(tuple(r), 1):
            ws.cell(row=i, column=c, value=v)
    for c, w in enumerate([13, 15, 15, 22, 24, 28, 15, 12, 11, 9, 14], 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{max(len(rows) + 1, 2)}"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ------------------------------------------------------------------ PDF
def attendance_pdf(report):
    need("reportlab", "reportlab")
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
                            title=f"Attendance Report {report['month']}")
    st = getSampleStyleSheet()
    t = report["totals"]
    story = [Paragraph(f"Attendance Report &ndash; {report['month']}", st["Title"]),
             Paragraph(f"Counted through {report['through']} &nbsp;|&nbsp; {t['employees']} employees &nbsp;|&nbsp; "
                       f"Present {t['present']} &nbsp; Late {t['late']} &nbsp; Absent {t['absent']} &nbsp; On leave {t['on_leave']} &nbsp;|&nbsp; "
                       f"Average attendance {t['average_rate'] if t['average_rate'] is not None else '-'}%", st["Normal"]),
             Spacer(1, 6 * mm)]
    data = [["Employee", "Department", "Shift", "Sched.", "Present", "Late", "Absent", "Leave", "Late min.", "Hours", "Rate"]]
    for r in report["rows"]:
        data.append([r["name"], r["department"], r["shift"], r["scheduled_days"], r["present"], r["late"], r["absent"],
                     r["on_leave"], r["late_minutes"], r["hours_worked"],
                     "-" if r["attendance_rate"] is None else f"{r['attendance_rate']}%"])
    tbl = Table(data, repeatRows=1, colWidths=[48 * mm, 38 * mm, 32 * mm] + [17 * mm] * 8)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F5EFF")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ALIGN", (3, 0), (-1, -1), "RIGHT"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5FA")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D7E2")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(tbl)
    doc.build(story)
    return out.getvalue()


def leave_form_pdf(db, leave_id):
    need("reportlab", "reportlab")
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    con = connect(db)
    r = con.execute("""
        SELECT lr.*, lt.name AS leave_type, e.first_name || ' ' || e.last_name AS emp_name, e.employee_no,
               COALESCE(d.name,'') AS department, COALESCE(e.position,'') AS position, rv.username AS reviewer
        FROM leave_requests lr JOIN leave_types lt ON lt.id = lr.leave_type_id JOIN employees e ON e.id = lr.employee_id
        LEFT JOIN departments d ON d.id = e.department_id LEFT JOIN users rv ON rv.id = lr.reviewed_by
        WHERE lr.id = ?""", (leave_id,)).fetchone()
    con.close()
    if not r:
        fail("Leave request not found.")

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm, bottomMargin=18 * mm,
                            title=f"Leave Request LR-{r['id']:05d}")
    st = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=st["Normal"], fontSize=9, textColor=colors.HexColor("#555555"))
    status_color = {"Approved": "#0E6A3C", "Rejected": "#A12B1F", "Pending": "#8A5A00", "Cancelled": "#4A5568"}.get(r["status"], "#000000")

    def row(label, value):
        return [Paragraph(f"<b>{label}</b>", st["Normal"]), Paragraph(str(value) if value not in (None, "") else "-", st["Normal"])]

    info = [row("Form No.", f"LR-{r['id']:05d}"), row("Employee", r["emp_name"]), row("Employee No.", r["employee_no"]),
            row("Department", r["department"]), row("Position", r["position"]), row("Leave type", r["leave_type"]),
            row("From", r["start_date"]), row("To", r["end_date"]), row("Working days", r["days"]),
            row("Reason", r["reason"]), row("Date filed", r["created_at"])]
    review = [row("Status", f"<font color='{status_color}'><b>{r['status'].upper()}</b></font>"),
              row("Reviewed by", r["reviewer"]), row("Reviewed on", r["reviewed_at"]), row("Remarks", r["review_remarks"])]
    style = TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D7E2")),
                        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)])
    t1 = Table(info, colWidths=[40 * mm, 130 * mm]); t1.setStyle(style)
    t2 = Table(review, colWidths=[40 * mm, 130 * mm]); t2.setStyle(style)
    sig = Table([["", ""], ["Employee signature / date", "HR Staff / Approver signature / date"]], colWidths=[85 * mm, 85 * mm], rowHeights=[18 * mm, None])
    sig.setStyle(TableStyle([("LINEABOVE", (0, 1), (-1, 1), 0.6, colors.black), ("FONTSIZE", (0, 1), (-1, 1), 8.5),
                             ("ALIGN", (0, 0), (-1, -1), "CENTER")]))
    doc.build([Paragraph("Leave Request Form", st["Title"]), Paragraph("Human Resource Information System", small), Spacer(1, 6 * mm),
               Paragraph("Request details", st["Heading3"]), t1, Spacer(1, 5 * mm),
               Paragraph("Decision", st["Heading3"]), t2, Spacer(1, 8 * mm), sig, Spacer(1, 10 * mm),
               Paragraph(f"Generated {datetime.now():%Y-%m-%d %H:%M} by the HRIS.", small)])
    return out.getvalue()


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("what", choices=["attendance", "leave-form", "employees"])
    ap.add_argument("--db", required=True)
    ap.add_argument("--month")
    ap.add_argument("--format", choices=["xlsx", "pdf"], default="xlsx")
    ap.add_argument("--today")
    ap.add_argument("--id", type=int)
    a = ap.parse_args()
    try:
        if a.what == "attendance":
            if not a.month:
                fail("--month is required")
            datetime.strptime(a.month, "%Y-%m")
            today = datetime.strptime(a.today, "%Y-%m-%d").date() if a.today else date.today()
            report = build_report(a.db, a.month, today)
            data = attendance_xlsx(report) if a.format == "xlsx" else attendance_pdf(report)
        elif a.what == "leave-form":
            if not a.id:
                fail("--id is required")
            data = leave_form_pdf(a.db, a.id)
        else:
            data = employees_xlsx(a.db)
    except SystemExit:
        raise
    except Exception as exc:
        fail(str(exc))
    sys.stdout.buffer.write(data)


if __name__ == "__main__":
    main()
